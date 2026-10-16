"""Load messy CSV and Excel exports into typed DataFrames.

Report-style exports (Gainskeeper, Silver, Aqua Data Studio result sets saved
from Excel) tend to have banner rows above the header, subtotal and footer
rows, page-break repeats of the header, accounting negatives like
``(1,234.56)``, and identifiers that lose their leading zeros the moment
anything parses them as numbers.

So every cell is read as text first, and a column is only converted when that
is safe: identifier-like columns always stay text.
"""

from __future__ import annotations

import csv
import io
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import numpy as np
import pandas as pd

DATA_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls"}

# Share of non-empty values that must parse before a column is converted.
PARSE_THRESHOLD = 0.95

# Column-name tokens that mark an identifier wherever they appear.
_STRONG_ID_TOKENS = {
    "cusip", "isin", "sedol", "figi", "ticker", "symbol", "id", "ssn", "tin",
    "taxid", "ein", "uuid", "guid",
}
# Tokens that make an identifier when the name is built only from them, or
# when one of the suffix tokens ends the name ("Account Number", "Trans Code").
_ID_NAME_TOKENS = {
    "account", "acct", "accountno", "acctno", "portfolio", "lot", "lotid",
    "security", "sec", "number", "no", "num", "nbr", "code", "id", "key",
}
_ID_SUFFIX_TOKENS = {"number", "no", "num", "nbr", "code", "id", "key"}

# Tokens that mark a column whose values should be masked in profiles.
_SENSITIVE_TOKENS = {
    "account", "acct", "accountno", "acctno", "portfolio", "ssn", "tin",
    "taxid", "ein", "name", "client", "owner", "holder", "beneficiary",
    "customer", "address", "addr", "street", "city", "zip", "postal",
    "email", "phone", "dob", "birth", "birthdate",
}

# "Security Name" or "Fund Name" describe holdings, not people.
_NOT_PERSONAL_TOKENS = {"security", "issuer", "fund", "asset", "product", "instrument", "file", "sheet", "column"}

_DATE_NAME_TOKENS = {"date", "dt", "asof", "settle", "settled", "trade", "acquired", "acq", "sold", "period", "month"}

_TOTAL_RE = re.compile(r"^\s*(grand\s+|sub\s*-?\s*)?totals?\b", re.I)
_NUMBER_RE = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_SCI_RE = re.compile(r"^\d(\.\d+)?[eE]\+\d+$")
_DATEISH_RE = re.compile(r"\d{1,4}[-/]\d{1,2}[-/]\d{1,4}")


@dataclass
class LoadedTable:
    """One sheet or CSV, cleaned, with a record of what was done to it."""

    name: str
    source: str
    sheet: str | None
    df: pd.DataFrame
    header_row: int  # 1-based row number in the source
    column_kinds: dict[str, str]  # identifier | numeric | datetime | text
    dropped: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# Column-name helpers
# --------------------------------------------------------------------------

def name_tokens(column: str) -> list[str]:
    """Split a column name into lowercase words: 'AcctNo_2' -> ['acct', 'no', '2']."""
    spaced = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", str(column))
    spaced = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", spaced)
    return [t for t in re.split(r"[^A-Za-z0-9]+", spaced.lower()) if t]


def is_identifier_name(column: str) -> bool:
    tokens = name_tokens(column)
    if not tokens:
        return False
    joined = "".join(tokens)
    if joined in _STRONG_ID_TOKENS or joined in _ID_NAME_TOKENS:
        return True
    if any(t in _STRONG_ID_TOKENS for t in tokens):
        return True
    if all(t in _ID_NAME_TOKENS for t in tokens):
        return True
    return tokens[-1] in _ID_SUFFIX_TOKENS and len(tokens) <= 3


def is_sensitive_name(column: str) -> bool:
    tokens = name_tokens(column)
    if set(tokens) & _NOT_PERSONAL_TOKENS:
        return False
    return any(t in _SENSITIVE_TOKENS for t in tokens) or "".join(tokens) in _SENSITIVE_TOKENS


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(text).lower()).strip("_") or "table"


# --------------------------------------------------------------------------
# Value parsing
# --------------------------------------------------------------------------

def parse_number(value: str) -> float | None:
    """Parse '1,234.56', '(1,234.56)', '$1,234', '12.5%', '1234.56-' -> float.

    Returns None when the text is not a number. A percent sign is stripped,
    not divided out: '12.5%' parses to 12.5.
    """
    s = str(value).strip()
    if not s:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    if s.endswith("-") and not s.startswith("-"):
        negative, s = True, s[:-1].strip()
    s = s.replace(",", "").replace("$", "").replace("%", "").replace(" ", "")
    if s.startswith("-"):
        negative, s = (not negative), s[1:]
    elif s.startswith("+"):
        s = s[1:]
    if not _NUMBER_RE.match(s):
        return None
    number = float(s)
    return -number if negative else number


def _cell_text(value, number_format: str | None = None) -> str:
    """Render an Excel cell value as the text a person would see."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        text = str(value)
        # A zero-padded format such as '000000000' means the leading zeros are
        # part of the identifier: put them back.
        if number_format and isinstance(value, int) and re.fullmatch(r"0+", number_format):
            text = text.zfill(len(number_format))
        return text
    return str(value).strip()


# --------------------------------------------------------------------------
# Raw readers: everything comes back as a list of rows of text
# --------------------------------------------------------------------------

def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _read_delimited(path: Path) -> list[list[str]]:
    text = _decode(path.read_bytes())
    sample = "\n".join(text.splitlines()[:50])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(path: Path) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[list[str]]] = {}
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows():
                rows.append([_cell_text(c.value, getattr(c, "number_format", None)) for c in row])
            sheets[ws.title] = rows
    finally:
        wb.close()
    return sheets


def _read_xls(path: Path) -> dict[str, list[list[str]]]:
    try:
        frames = pd.read_excel(path, sheet_name=None, header=None, dtype=str, engine="xlrd")
    except ImportError as exc:
        raise RuntimeError(
            f"{path.name} is a legacy .xls file and the 'xlrd' package is not installed. "
            "Install it (pip install xlrd) or re-save the file as .xlsx."
        ) from exc
    return {
        name: [["" if pd.isna(v) else str(v).strip() for v in row] for row in frame.itertuples(index=False)]
        for name, frame in frames.items()
    }


# --------------------------------------------------------------------------
# Structure: header row, junk rows
# --------------------------------------------------------------------------

def _nonempty(row: list[str]) -> int:
    return sum(1 for c in row if c)


def _looks_like_value(cell: str) -> bool:
    return parse_number(cell) is not None or bool(_DATEISH_RE.search(cell))


def find_header_row(rows: list[list[str]], scan: int = 30) -> int:
    """Index of the first row that looks like column headers.

    A header row is nearly as wide as the widest row near the top of the file
    and is made of labels, not numbers or dates.
    """
    window = rows[:scan]
    widths = [_nonempty(r) for r in window]
    if not widths or max(widths) == 0:
        return 0
    target = max(widths)
    if target < 2:
        return next(i for i, w in enumerate(widths) if w)
    for i, row in enumerate(window):
        cells = [c for c in row if c]
        if len(cells) < max(2, 0.6 * target):
            continue
        labels = sum(1 for c in cells if not _looks_like_value(c))
        if labels / len(cells) >= 0.8:
            return i
    return next(i for i, w in enumerate(widths) if w)


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, name in enumerate(names):
        name = re.sub(r"\s+", " ", name).strip() or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        out.append(name)
    return out


def _is_total_row(row: list[str]) -> bool:
    leading = [c for c in row if c][:2]
    return any(len(c) <= 60 and _TOTAL_RE.match(c) for c in leading)


def _structure(rows: list[list[str]]) -> tuple[int, list[str], list[list[str]], list[dict]]:
    """Find the header and drop blank, total, repeated-header and footer rows."""
    header_idx = find_header_row(rows)
    width = max((len(r) for r in rows), default=0)
    padded = [r + [""] * (width - len(r)) for r in rows]
    raw_header = padded[header_idx] if padded else []

    # Drop columns that are empty in the header and every data row.
    keep = [j for j in range(width) if any(r[j] for r in padded[header_idx:])]
    header = _dedupe([raw_header[j] for j in keep])
    header_cells = [raw_header[j] for j in keep]

    body: list[tuple[int, list[str]]] = []
    dropped: list[dict] = []
    for offset, row in enumerate(padded[header_idx + 1:]):
        source_row = header_idx + 2 + offset  # 1-based
        cells = [row[j] for j in keep]
        if not any(cells):
            continue
        if cells == header_cells:
            dropped.append({"row": source_row, "reason": "repeated header", "preview": _preview(cells)})
            continue
        if _is_total_row(cells):
            dropped.append({"row": source_row, "reason": "total/subtotal row", "preview": _preview(cells)})
            continue
        body.append((source_row, cells))

    # Footer: trailing rows with a single filled cell ("Report generated ...").
    if len(header) > 1:
        while body and _nonempty(body[-1][1]) == 1:
            source_row, cells = body.pop()
            dropped.append({"row": source_row, "reason": "footer row", "preview": _preview(cells)})

    dropped.sort(key=lambda d: d["row"])
    return header_idx + 1, header, [cells for _, cells in body], dropped


def _preview(cells: list[str], limit: int = 80) -> str:
    text = " | ".join(c for c in cells if c)
    return text if len(text) <= limit else text[: limit - 1] + "…"


# --------------------------------------------------------------------------
# Typing
# --------------------------------------------------------------------------

def _classify_and_convert(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, str], list[str]]:
    kinds: dict[str, str] = {}
    warnings: list[str] = []
    out = {}
    for col in df.columns:
        values = df[col].fillna("").astype(str).str.strip()
        filled = values[values != ""]
        if is_identifier_name(col) or _has_leading_zero_digits(filled):
            cleaned = values.where(~values.str.fullmatch(r"\d+\.0"), values.str[:-2])
            if filled.str.fullmatch(_SCI_RE.pattern).any():
                warnings.append(
                    f"'{col}' has values in scientific notation (e.g. 1.2E+10): the source "
                    "was opened in Excel as a number and digits were lost. Re-export as text."
                )
            variants = filled.drop_duplicates()
            clashes = variants.groupby(variants.str.lstrip("0")).filter(lambda g: len(g) > 1)
            if not clashes.empty:
                clashes = clashes.iloc[clashes.str.lstrip("0").argsort(kind="stable")]
                example = clashes.head(2).tolist()
                warnings.append(
                    f"'{col}': some values differ only by leading zeros (e.g. {' vs '.join(example)}). "
                    "Normalize before joining or reconciling."
                )
            if "cusip" in name_tokens(col) and filled.size:
                bad = int((filled.str.len() != 9).sum())
                if bad:
                    warnings.append(f"'{col}': {bad} value(s) are not 9 characters; leading zeros may be missing.")
            out[col] = _blank_to_na(cleaned)
            kinds[col] = "identifier"
            continue

        if filled.empty:
            out[col] = _blank_to_na(values)
            kinds[col] = "text"
            continue

        numbers = filled.map(parse_number)
        if numbers.notna().mean() >= PARSE_THRESHOLD:
            out[col] = values.map(parse_number).astype("float64")
            kinds[col] = "numeric"
            unparsed = int(numbers.isna().sum())
            if unparsed:
                warnings.append(f"'{col}': {unparsed} non-numeric value(s) treated as missing.")
            continue

        name_says_date = bool(set(name_tokens(col)) & _DATE_NAME_TOKENS)
        if name_says_date or filled.str.contains(_DATEISH_RE).mean() >= PARSE_THRESHOLD:
            parsed = _to_datetime(values)
            if parsed[values != ""].notna().mean() >= PARSE_THRESHOLD:
                out[col] = parsed
                kinds[col] = "datetime"
                continue

        out[col] = _blank_to_na(values)
        kinds[col] = "text"
    return pd.DataFrame(out, index=df.index), kinds, warnings


def _blank_to_na(values: pd.Series) -> pd.Series:
    return values.mask(values == "")


def _has_leading_zero_digits(filled: pd.Series) -> bool:
    if filled.empty:
        return False
    digits = filled.str.fullmatch(r"\d+")
    return bool(digits.all() and filled.str.match(r"0\d").any())


def _to_datetime(values: pd.Series) -> pd.Series:
    cleaned = _blank_to_na(values)
    try:
        return pd.to_datetime(cleaned, errors="coerce", format="mixed")
    except (TypeError, ValueError):
        return pd.to_datetime(cleaned, errors="coerce")


# --------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------

def table_from_rows(rows: list[list[str]], name: str, source: str, sheet: str | None = None) -> LoadedTable | None:
    if not any(_nonempty(r) for r in rows):
        return None
    header_row, header, body, dropped = _structure(rows)
    raw = pd.DataFrame(body, columns=header) if body else pd.DataFrame(columns=header)
    df, kinds, warnings = _classify_and_convert(raw)
    df = df.reset_index(drop=True)
    if header_row > 1:
        warnings.insert(0, f"Header found on row {header_row}; {header_row - 1} row(s) above it skipped.")
    return LoadedTable(
        name=name, source=source, sheet=sheet, df=df, header_row=header_row,
        column_kinds=kinds, dropped=dropped, warnings=warnings,
    )


def load_file(path: str | Path) -> list[LoadedTable]:
    """Load every non-empty table in a CSV or workbook."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in {".csv", ".tsv", ".txt"}:
        table = table_from_rows(_read_delimited(path), slug(path.stem), path.name)
        return [table] if table else []
    if ext in {".xlsx", ".xlsm"}:
        sheets = _read_xlsx(path)
    elif ext == ".xls":
        sheets = _read_xls(path)
    else:
        raise ValueError(f"Unsupported file type: {path.name}")
    tables = []
    for sheet, rows in sheets.items():
        name = slug(path.stem) if len(sheets) == 1 else f"{slug(path.stem)}__{slug(sheet)}"
        table = table_from_rows(rows, name, path.name, sheet)
        if table is not None and not table.df.empty:
            tables.append(table)
    return tables


def load_table(path: str | Path, sheet: str | None = None) -> pd.DataFrame:
    """Load one CSV, or one sheet of a workbook (the first non-empty one by default)."""
    tables = load_file(path)
    if not tables:
        raise ValueError(f"No data found in {path}")
    if sheet is None:
        return tables[0].df
    for t in tables:
        if t.sheet == sheet:
            return t.df
    raise ValueError(f"Sheet '{sheet}' not found in {path}; have {[t.sheet for t in tables]}")


def data_files(workspace: str | Path = "workspace") -> list[Path]:
    data_dir = Path(workspace) / "data"
    if not data_dir.is_dir():
        return []
    return sorted(
        p for p in data_dir.iterdir()
        if p.is_file() and p.suffix.lower() in DATA_EXTENSIONS and not p.name.startswith(("~$", "."))
    )


def load_all(workspace: str | Path = "workspace") -> dict[str, LoadedTable]:
    """Load every table under <workspace>/data, keyed by table name."""
    tables: dict[str, LoadedTable] = {}
    for path in data_files(workspace):
        for table in load_file(path):
            name, n = table.name, 1
            while name in tables:
                n += 1
                name = f"{table.name}_{n}"
            table.name = name
            tables[name] = table
    return tables


def safe_stdout() -> None:
    """Don't crash printing a value the console can't encode (legacy Windows code pages)."""
    reconfigure = getattr(sys.stdout, "reconfigure", None)
    if reconfigure:
        reconfigure(errors="replace")


def numeric_values(df: pd.DataFrame) -> np.ndarray:
    """Every finite numeric cell in a frame, flattened (used by the grounding check)."""
    cols = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c])]
    if not cols:
        return np.array([], dtype=float)
    arr = df[cols].to_numpy(dtype=float, na_value=np.nan).ravel()
    return arr[np.isfinite(arr)]
