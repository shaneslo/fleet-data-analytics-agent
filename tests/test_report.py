import json

import pandas as pd
from openpyxl import load_workbook

from analyst import chart, profile, report
from analyst.run import latest_run, new_run


def _run_with_tables(workspace):
    profile.main(["--workspace", str(workspace)])
    run = new_run("Why do GK and Silver disagree?", workspace)
    pd.DataFrame({"Account": ["00001234", "00005678"], "breaks": [3, 1], "net_difference": [-1250.5, 4.0]}).to_csv(
        run / "analysis" / "breaks_by_account.csv", index=False)
    assert chart.main(["--run", str(run), "--data", "analysis/breaks_by_account.csv", "--type", "barh",
                       "--x", "Account", "--y", "breaks", "--output", "charts/breaks.png"]) == 0
    return run


def _findings(run, answer, detail="Net difference of $1,250.50 sits in one account."):
    (run / "findings.json").write_text(json.dumps({
        "title": "GK vs Silver breaks",
        "answer": answer,
        "findings": [{"headline": "Account 00001234 holds 3 of the 4 breaks", "detail": detail,
                      "table": "analysis/breaks_by_account.csv", "chart": "charts/breaks.png"}],
        "caveats": ["Silver export is as of trade date."],
        "next_steps": ["Pull the Sedona ledger for 00001234."],
    }))


def test_report_grounded(workspace):
    run = _run_with_tables(workspace)
    _findings(run, "Breaks total 4 across 2 accounts; GK has 3 rows and Silver 3 rows.")
    assert report.main(["--run", str(run), "--workspace", str(workspace)]) == 0
    md = (run / "report.md").read_text(encoding="utf-8")
    assert "## Answer" in md and "![" in md and "Unverified figures" not in md
    wb = load_workbook(run / "report.xlsx")
    assert wb.sheetnames[:2] == ["Summary", "Data used"]
    sheet = wb["breaks_by_account"]
    assert sheet["A2"].value == "00001234" and sheet["A2"].number_format == "@"
    assert latest_run(workspace) == run


def test_report_flags_invented_figure(workspace, capsys):
    run = _run_with_tables(workspace)
    _findings(run, "Silver is short by $98,765.43 overall.")
    assert report.main(["--run", str(run), "--workspace", str(workspace)]) == 0
    assert "Unverified figures" in (run / "report.md").read_text(encoding="utf-8")
    assert "$98,765.43" in capsys.readouterr().out
    assert report.main(["--run", str(run), "--workspace", str(workspace), "--strict"]) == 1


def test_rounded_and_scaled_figures_match(workspace):
    run = _run_with_tables(workspace)
    _findings(run, "About $1.3K of net difference, 75% of breaks in one account.",
              detail="Net difference of $1,250.5.")
    tables = report.load_analysis(run)
    (run / "analysis" / "share.csv").write_text("share\n0.75\n")
    tables = report.load_analysis(run)
    values, ids = report.known_values(tables, None)
    findings = report.load_findings(run)
    assert report.unverified_figures(findings, values, ids) == []


def test_missing_reference_is_an_error(workspace):
    run = _run_with_tables(workspace)
    (run / "findings.json").write_text(json.dumps({
        "title": "t", "answer": "a", "findings": [{"headline": "h", "chart": "charts/nope.png"}]}))
    assert report.main(["--run", str(run), "--workspace", str(workspace)]) == 1


def test_tables_only_without_findings(workspace, capsys):
    run = _run_with_tables(workspace)
    assert report.main(["--run", str(run), "--workspace", str(workspace)]) == 0
    assert "No findings.json" in (run / "report.md").read_text(encoding="utf-8")
