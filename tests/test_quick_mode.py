import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from analyst import profile
from analyst.xlsx import export, infer_kinds

SKILL = Path(__file__).resolve().parents[1] / ".github" / "skills" / "quick-answer" / "SKILL.md"


def test_export_keeps_ids_text_and_formulas_inert(tmp_path):
    df = pd.DataFrame({"Account": ["00001234"], "note": ["=1+1"], "Qty": [5.0],
                       "Trade Date": pd.to_datetime(["2026-09-01"])})
    path = export(df, "Q3 losses", workspace=tmp_path)
    assert path.parent == tmp_path / "exports" and path.name.startswith("q3_losses_")
    ws = load_workbook(path).active
    assert ws["A2"].value == "00001234" and ws["A2"].number_format == "@"
    assert ws["B2"].data_type == "s" and ws["B2"].value == "=1+1"
    assert ws["C2"].number_format == "#,##0" and ws["D2"].number_format == "yyyy-mm-dd"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_infer_kinds():
    df = pd.DataFrame({"AcctNo": ["1"], "Desc": ["x"], "Qty": [1.0], "When": pd.to_datetime(["2026-01-01"])})
    assert infer_kinds(df) == {"AcctNo": "identifier", "Desc": "text", "Qty": "numeric", "When": "datetime"}


def test_profile_if_changed_reuses_then_refreshes(workspace, capsys):
    assert profile.main(["--workspace", str(workspace)]) == 0
    out = workspace / "profile.json"
    first = out.read_text(encoding="utf-8")
    capsys.readouterr()

    assert profile.main(["--workspace", str(workspace), "--if-changed"]) == 0
    assert "Profile is current" in capsys.readouterr().out
    assert out.read_text(encoding="utf-8") == first

    extra = workspace / "data" / "new.csv"
    extra.write_text("Account,Qty\n00000001,1\n", encoding="utf-8")
    assert profile.main(["--workspace", str(workspace), "--if-changed"]) == 0
    assert "Profile is current" not in capsys.readouterr().out
    assert "new" in {t["name"] for t in json.loads(out.read_text(encoding="utf-8"))["tables"]}


def test_skill_example_runs_on_fixture(workspace, tmp_path):
    """The quick-answer skill's example script must run as written."""
    blocks = re.findall(r"```python\n(.*?)```", SKILL.read_text(encoding="utf-8"), re.DOTALL)
    script = tmp_path / "scratch.py"
    script.write_text(blocks[0], encoding="utf-8")
    result = subprocess.run([sys.executable, str(script)], cwd=workspace.parent, capture_output=True, text=True,
                            env={**os.environ, "PYTHONPATH": str(Path(__file__).resolve().parents[1])})
    assert result.returncode == 0, result.stderr
    assert "realized loss total: -45.1" in result.stdout


def test_numeric_account_column_is_written_as_text(tmp_path):
    df = pd.DataFrame({"Account": [1234567890123456789, 42], "Account Value": [5000.0, 1.0]})
    ws = load_workbook(export(df, "ids", workspace=tmp_path)).active
    assert ws["A2"].value == "1234567890123456789" and ws["A2"].data_type == "s"
    assert ws["A3"].value == "42"
    assert ws["B2"].value == 5000  # a measure stays numeric


def test_oversized_frames_split_across_sheets(tmp_path, monkeypatch, capsys):
    from analyst import xlsx
    monkeypatch.setattr(xlsx, "EXCEL_MAX_ROWS", 3)
    df = pd.DataFrame({"Qty": range(7)})
    wb = load_workbook(export(df, "big", workspace=tmp_path))
    assert wb.sheetnames == ["big", "big_2", "big_3"]
    assert sum(ws.max_row - 1 for ws in wb.worksheets) == 7
    assert "7 row(s) across 3 sheets" in capsys.readouterr().out


def test_if_changed_sees_same_size_rewrite_within_a_second(workspace, capsys):
    profile.main(["--workspace", str(workspace)])
    path = workspace / "data" / "silver_positions.csv"
    stat = path.stat()
    path.write_bytes(path.read_bytes().replace(b"Alpha Corp", b"Alpha Co.."))  # same size
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000))  # same second, 1 microsecond later
    capsys.readouterr()
    profile.main(["--workspace", str(workspace), "--if-changed"])
    assert "Profile is current" not in capsys.readouterr().out
